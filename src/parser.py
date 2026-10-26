"""Resume / job-description parsing: text extraction, section segmentation and
structured information extraction (contact details, education, experience)."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date

# Canonical section name -> headings that introduce it in a resume.
SECTION_HEADINGS = {
    "summary": ["summary", "profile", "professional summary", "career objective", "objective",
                "about me", "career summary", "personal statement"],
    "skills": ["skills", "technical skills", "key skills", "core competencies", "competencies",
               "skill set", "technologies", "tools and technologies", "technical expertise", "expertise"],
    "experience": ["experience", "work experience", "professional experience", "employment history",
                   "work history", "employment", "career history", "internships", "internship"],
    "education": ["education", "academic background", "academics", "qualifications",
                  "educational qualifications", "academic qualifications", "education and training"],
    "projects": ["projects", "academic projects", "personal projects", "key projects", "project experience"],
    "certifications": ["certifications", "certificates", "licenses and certifications", "courses",
                       "trainings", "training"],
    "achievements": ["achievements", "awards", "honors", "accomplishments", "awards and achievements"],
    "languages": ["languages"],
    "interests": ["interests", "hobbies"],
}
_HEADING_LOOKUP = {h: sec for sec, hs in SECTION_HEADINGS.items() for h in hs}

DEGREE_PATTERNS = [
    (r"\b(ph\.?\s?d|doctorate)\b", "PhD"),
    (r"\b(m\.?\s?s\.?c?|master'?s?|m\.?\s?tech|m\.?\s?e\.|mba|m\.?\s?phil|mcs)\b", "Master's"),
    (r"\b(b\.?\s?s\.?c?|bachelor'?s?|b\.?\s?tech|b\.?\s?e\.|bba|bcs|b\.?\s?com|undergraduate)\b", "Bachelor's"),
    (r"\b(associate degree|diploma)\b", "Diploma"),
    (r"\b(intermediate|hssc|f\.?\s?sc|a[- ]levels?|high school)\b", "Intermediate/High School"),
]
DEGREE_RANK = {"PhD": 5, "Master's": 4, "Bachelor's": 3, "Diploma": 2, "Intermediate/High School": 1}

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}
_MONTH_RE = r"(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
_DATE_RE = rf"(?:{_MONTH_RE}\s*,?\s*)?((?:19|20)\d{{2}})|(?:(\d{{1,2}})/((?:19|20)\d{{2}}))"
DATE_RANGE_RE = re.compile(
    rf"(?P<start>{_DATE_RE})\s*(?:-|–|—|to|till|until)\s*(?P<end>{_DATE_RE}|present|current|now|date|ongoing)",
    re.IGNORECASE,
)
EXPLICIT_YEARS_RE = re.compile(
    r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)\s*(?:of\s+)?(?:\w+\s+){0,3}?experience", re.IGNORECASE)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{3,4}\)?[\s-]?)\d{3}[\s-]?\d{3,4}")
URL_RE = re.compile(r"(?:https?://|www\.)\S+|(?:linkedin\.com|github\.com)/\S+", re.IGNORECASE)


@dataclass
class ResumeProfile:
    raw_text: str
    name: str = ""
    email: str = ""
    phone: str = ""
    links: list[str] = field(default_factory=list)
    sections: dict[str, str] = field(default_factory=dict)
    education: list[dict] = field(default_factory=list)
    highest_degree: str = ""
    experience_years: float = 0.0
    skills: list[str] = field(default_factory=list)
    skills_by_category: dict[str, list[str]] = field(default_factory=dict)


# --------------------------------------------------------------------------- text extraction

def extract_text(data: bytes, filename: str) -> str:
    """Extract plain text from a PDF, DOCX or TXT file given as raw bytes."""
    name = filename.lower()
    if name.endswith(".pdf"):
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
        return normalize_whitespace("\n".join(pages))
    if name.endswith(".docx"):
        import docx
        document = docx.Document(io.BytesIO(data))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return normalize_whitespace("\n".join(lines))
    if name.endswith(".txt") or name.endswith(".md"):
        return normalize_whitespace(data.decode("utf-8", errors="ignore"))
    raise ValueError(f"Unsupported file type: {filename} (use PDF, DOCX or TXT)")


def normalize_whitespace(text: str) -> str:
    text = text.replace("\r", "\n").replace("•", "\n• ").replace("", "\n• ")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


# --------------------------------------------------------------------------- sections

def _heading_of(line: str) -> str | None:
    cleaned = re.sub(r"[^a-z& ]", "", line.lower()).replace("&", "and").strip()
    if not cleaned or len(cleaned.split()) > 5:
        return None
    return _HEADING_LOOKUP.get(cleaned)


def split_sections(text: str) -> dict[str, str]:
    """Split resume text into canonical sections. Text before the first heading is 'header'."""
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for line in text.splitlines():
        stripped = line.strip()
        heading = _heading_of(stripped.rstrip(":"))
        if heading is None and ":" in stripped:
            # "Skills: Python, SQL" style inline heading
            head, _, rest = stripped.partition(":")
            inline = _heading_of(head)
            if inline:
                current = inline
                sections.setdefault(current, []).append(rest.strip())
                continue
        if heading:
            current = heading
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(stripped)
    return {k: "\n".join(v).strip() for k, v in sections.items() if "\n".join(v).strip()}


# --------------------------------------------------------------------------- field extraction

def extract_contact(text: str) -> tuple[str, str, str, list[str]]:
    email = (EMAIL_RE.search(text) or [""])[0]
    phone = ""
    for m in PHONE_RE.finditer(text):
        digits = re.sub(r"\D", "", m.group())
        if 10 <= len(digits) <= 15 and not re.fullmatch(r"(19|20)\d{2}(19|20)\d{2}", digits):
            phone = m.group().strip()
            break
    links = sorted(set(URL_RE.findall(text)))
    name = ""
    for line in text.splitlines()[:6]:
        line = line.strip()
        words = line.split()
        if (1 < len(words) <= 4 and not EMAIL_RE.search(line) and not any(c.isdigit() for c in line)
                and _heading_of(line) is None and all(w[0].isupper() for w in words if w[0].isalpha())):
            name = line
            break
    return name, email, phone, links


def extract_education(text: str, sections: dict[str, str]) -> list[dict]:
    source = sections.get("education") or text
    entries = []
    for line in source.splitlines():
        lowered = line.lower()
        for pattern, degree in DEGREE_PATTERNS:
            if re.search(pattern, lowered):
                year = re.findall(r"(?:19|20)\d{2}", line)
                entries.append({"degree": degree, "details": line.strip(" •-"),
                                "year": year[-1] if year else ""})
                break
    return entries


def _to_months(year: str | None, month: str | None, numeric_month: str | None) -> int:
    m = 1
    if month:
        m = MONTHS.get(month.lower()[:3], 1)
    elif numeric_month:
        m = max(1, min(12, int(numeric_month)))
    return int(year) * 12 + m


def _parse_point(fragment: str, is_end: bool) -> int | None:
    fragment = fragment.strip().lower()
    if fragment in {"present", "current", "now", "date", "ongoing"}:
        today = date.today()
        return today.year * 12 + today.month
    m = re.search(_DATE_RE, fragment, re.IGNORECASE)
    if not m:
        return None
    month, year, num_month, num_year = m.group(1), m.group(2), m.group(3), m.group(4)
    if year:
        months = _to_months(year, month, None)
        if is_end and not month:
            months += 11  # "2019 - 2021" counts the whole final year
        return months
    return _to_months(num_year, None, num_month)


def estimate_experience_years(text: str, sections: dict[str, str]) -> float:
    """Estimate total years of experience from date ranges in the experience section
    (overlapping ranges are merged); fall back to an explicit 'N years of experience'."""
    source = sections.get("experience", "")
    intervals = []
    for m in DATE_RANGE_RE.finditer(source):
        start = _parse_point(m.group("start"), is_end=False)
        end = _parse_point(m.group("end"), is_end=True)
        if start and end and end >= start:
            intervals.append((start, end))
    total_months = 0
    if intervals:
        intervals.sort()
        cur_s, cur_e = intervals[0]
        for s, e in intervals[1:]:
            if s <= cur_e:
                cur_e = max(cur_e, e)
            else:
                total_months += cur_e - cur_s
                cur_s, cur_e = s, e
        total_months += cur_e - cur_s
    from_dates = round(total_months / 12, 1)
    explicit = [float(x) for x in EXPLICIT_YEARS_RE.findall(text)]
    if from_dates == 0 and explicit:
        return max(explicit)
    return from_dates


def parse_resume(text: str) -> ResumeProfile:
    text = normalize_whitespace(text)
    sections = split_sections(text)
    name, email, phone, links = extract_contact(text)
    education = extract_education(text, sections)
    highest = max((e["degree"] for e in education), key=lambda d: DEGREE_RANK[d], default="")
    return ResumeProfile(
        raw_text=text, name=name, email=email, phone=phone, links=links, sections=sections,
        education=education, highest_degree=highest,
        experience_years=estimate_experience_years(text, sections),
    )


# --------------------------------------------------------------------------- job descriptions

PREFERRED_MARKERS = re.compile(
    r"\b(preferred|nice[- ]to[- ]have|good[- ]to[- ]have|bonus|a plus|is a plus|desired|desirable|"
    r"advantage|advantageous|optional|familiarity with .* (?:a )?plus)\b", re.IGNORECASE)
PREFERRED_HEADINGS = re.compile(
    r"^\s*(preferred|nice to have|good to have|bonus|desired|additional|plus)\b.*:?\s*$", re.IGNORECASE)
REQUIRED_HEADINGS = re.compile(
    r"^\s*(requirements?|required|qualifications|must[- ]have|what you need|what we.re looking for|"
    r"skills|responsibilities|key responsibilities|about the role|minimum qualifications)\b.*:?\s*$",
    re.IGNORECASE)


def split_job_requirements(description: str) -> tuple[str, str]:
    """Split a job description into (required_text, preferred_text).
    A line is 'preferred' if it sits under a preferred-style heading or itself
    contains a marker like 'preferred', 'nice to have', 'is a plus'."""
    required, preferred = [], []
    in_preferred = False
    for line in normalize_whitespace(description).splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if len(stripped.split()) <= 6 and PREFERRED_HEADINGS.match(stripped):
            in_preferred = True
            continue
        if len(stripped.split()) <= 6 and REQUIRED_HEADINGS.match(stripped):
            in_preferred = False
            continue
        if in_preferred or PREFERRED_MARKERS.search(stripped):
            preferred.append(stripped)
        else:
            required.append(stripped)
    return "\n".join(required), "\n".join(preferred)


def required_experience_years(description: str) -> float:
    """Minimum years of experience a job asks for (0 if not stated)."""
    values = re.findall(r"(\d{1,2})\s*\+?\s*(?:-\s*\d{1,2}\s*)?(?:years?|yrs?)", description, re.IGNORECASE)
    return float(min(int(v) for v in values)) if values else 0.0
